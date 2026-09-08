import difflib
import sys
import unicodedata
from datetime import datetime
from io import BytesIO
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from leitura_csv import ler_csv_flexivel

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(errors="replace")

# ==========================================================================
# Configuração / constantes
# ==========================================================================

BASE_DIR = Path(__file__).resolve().parent
PASTA_INPUT = BASE_DIR / "input"
PASTA_OUTPUT = BASE_DIR / "output"
PASTA_EXPORT = PASTA_OUTPUT / "export"

COLUNAS_OBRIGATORIAS = [
    'Code Deep',
    'Nome do Agente',
    'Status_Validacao',
    'Status_Aprovacao_Backoffice',
]

COLUNAS_SITUACAO_BACKOFFICE = [
    'Situação Backoffice do Bloco',
    'Situação Backoffice',
]

COLUNAS_ALGUEM_ENCONTRADO = [
    'Alguém foi encontrado?_field',
    '1.2 Alguém foi encontrado?',
]
COLUNAS_MUNICIPIO = [
    '1.4 Município_field',
    '1.4 Município',
]
COLUNAS_CODIGO_UNICO_CLIENTE = [
    'codigo_unico_cliente',
]

COLUNAS_VAZIAS_BACKOFFICE = [
    'Situação.1', 'Data de conferência', 'Conferência', 'Responsável',
    'Data de reconferência', 'Reconferência', 'Responsável_Reconf',
    'Disponível p/ conferência', 'Disponível p/ reconferência',
]

ORDEM_FINAL_BASE = [
    'Link Backoffice', 'Code Deep', 'CODE', 'Nome do Agente', 'Data do registro', 'Situação',
    'Situação.1', 'Data de conferência', 'Conferência', 'Responsável',
    'Data de reconferência', 'Reconferência', 'Responsável_Reconf',
    'Disponível p/ conferência', 'Disponível p/ reconferência',
    'Tipo_Visita', 'Indicativo_Visita', '1.2 Alguém foi encontrado?',
    'Alguém foi encontrado?_field', '1.4 Município_field', 'codigo_unico_cliente',
    'codigo_unico', 'Sinalização Código', 'Detalhe_Inconsistencia', 'Erros_Logicos',
]

COLUNAS_OCULTAS_BASE = [
    'Reconferência',
    'Responsável_Reconf',
    'Disponível p/ conferência',
    'Disponível p/ reconferência',
]

LETRAS_ESPERADAS_BASE = {
    'Situação.1': 'G',
    'Data de conferência': 'H',
    'Conferência': 'I',
    'Responsável': 'J',
    'Code Deep': 'B',
    'CODE': 'C',
    'Nome do Agente': 'D',
}

OPCOES_SITUACAO_1 = ['APROVADO', 'REPROVADO', 'AGUARDANDO AJUSTE']
OPCOES_CONFERENCIA = ['FEITO', 'ENVIAR P AGENTE']

LIMIAR_SIMILARIDADE_MUNICIPIO = 0.8
ALIASES_MUNICIPIOS = {}

# ==========================================================================
# LISTA DE RESPONSÁVEIS ATUALIZADA (COM O NOME QUE FALTAVA)
# ==========================================================================
RESPONSAVEIS_BACKOFFICE = [
    'Alessandra Peixoto Fidelis',
    'Andre de Oliveira Bertole',
    'Caio Viapiana',
    'Debora Carvalho Cavalcanti',
    'Denis Takeshi Sakuma',
    'Diego Gonçalves dos Santos',
    'Douglas Jordão Moreira',
    'Felipe Araujo',
    'Fernanda Luara',
    'Giovana Carvalho',
    'Guilherme Araujo Brito da Silva',
    'Guilherme Collin Dias',
    'Herysson Alves Pereira',
    'Hudson Bento Francelino',
    'Hudson Salmistraro',
    'Icaro Ryan',
    'Isabella Freitas Pinto',
    'Jhonata Soares',
    'José Artur da Costa',
    'José Bernardo Moreira Pinto de Almeida',
    'Julio Cesar da Silva Matos',
    'Kayky Souza',
    'Kelson Henrique dos Santos Pereira',
    'Keren Gonçalves Ferreira',
    'Letícia de Campos',
    'Lorrayne Calazans Braga',
    'Lucas dos Santos Martins',
    'Luiz Henrique Soares Lobato',
    'Marcelo Hideki Ogushi Carrera',
    'Maria Eduarda Amancio Silva',
    'Mathesu Francisco da Silva Barbosa',   # <-- NOME ADICIONADO
    'Miguel Archanjo Shigueru Kimura',
    'Pedro Henrique dos Santos Firmino',
    'Pedro Henrique Nori Takara',
    'Rafael Rodrigues de Almeida',
    'Rafael Silva',
    'Viviana Aparecida da Silva',
    'Wesley Xavier',
    'Winicius Khaue Reis e Silva',
]

ALVOS_DESMEMBRAMENTO = 'Desmembramento|Fraude|Duplicata|Duplicidade|Parecida'
ALVOS_EXCLUSAO = 'Desmembramento|Fraude'
COLUNAS_DESMEMBRADAS = [
    'Link Backoffice', 'Code Deep', 'Situação', 'Data de conferência',
    'Conferência', 'Responsável',
]

BORDA_FINA = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
FUNDO_HEADER_AZUL_ESCURO = PatternFill("solid", fgColor="1F4E78")
FUNDO_HEADER_AZUL_CLARO = PatternFill("solid", fgColor="BDD7EE")
FUNDO_PAINEL_VERMELHO = PatternFill("solid", fgColor="C00000")
FUNDO_TOTAL = PatternFill("solid", fgColor="D9D9D9")
FONTE_HEADER_BRANCA = Font(bold=True, color="FFFFFF")
FONTE_HEADER_PRETA = Font(bold=True)
CENTRO = Alignment(horizontal="center")
ESQUERDA = Alignment(horizontal="left")


class ErroProcessamento(Exception):
    pass


# ==========================================================================
# Etapa 1 — Município
# ==========================================================================

def normalizar_municipio(nome: str) -> str:
    return nome.strip().upper()


def normalizar_texto(texto) -> str:
    texto = str(texto).strip().upper()
    texto = unicodedata.normalize('NFD', texto)
    texto = ''.join(c for c in texto if not unicodedata.combining(c))
    return ' '.join(texto.split())


def solicitar_municipio() -> str:
    print("=" * 40)
    print("CONFERÊNCIA BACKOFFICE")
    print("=" * 40)
    print()
    nome = input("Digite o nome do município: ")
    municipio = normalizar_municipio(nome)
    if not municipio:
        raise ErroProcessamento("ERRO: nome de município inválido.")
    return municipio


# ==========================================================================
# Etapa 2 — Arquivo de entrada
# ==========================================================================

def obter_arquivo_input(pasta_input: Path) -> Path:
    pasta_input.mkdir(parents=True, exist_ok=True)
    arquivos = sorted(pasta_input.glob("*.xlsx"))
    if not arquivos:
        raise ErroProcessamento("Nenhum arquivo .xlsx encontrado na pasta input.")
    if len(arquivos) == 1:
        return arquivos[0]
    print(f"Foram encontrados {len(arquivos)} arquivos:\n")
    for i, arquivo in enumerate(arquivos, start=1):
        print(f"{i} - {arquivo.name}")
    print()
    while True:
        escolha = input("Digite o número do arquivo que deseja processar: ").strip()
        if escolha.isdigit() and 1 <= int(escolha) <= len(arquivos):
            return arquivos[int(escolha) - 1]
        print("Opção inválida. Tente novamente.")


# ==========================================================================
# Etapa 3 — Validação e carga da base
# ==========================================================================

def validar_colunas(df: pd.DataFrame, colunas_obrigatorias: list) -> None:
    faltantes = [c for c in colunas_obrigatorias if c not in df.columns]
    if faltantes:
        for coluna in faltantes:
            print(f"ERRO: coluna obrigatória não encontrada: {coluna}")
        raise ErroProcessamento("Processamento interrompido por colunas obrigatórias ausentes.")


def carregar_base(caminho_arquivo: Path) -> pd.DataFrame:
    eh_excel = False
    conteudo = None
    if hasattr(caminho_arquivo, 'read'):
        posicao = caminho_arquivo.tell()
        assinatura = caminho_arquivo.read(4)
        caminho_arquivo.seek(posicao)
        eh_excel = assinatura.startswith(b'PK') or assinatura == b'\xd0\xcf\x11\xe0'
        if not eh_excel:
            conteudo = caminho_arquivo.read()
            caminho_arquivo.seek(posicao)
    else:
        eh_excel = str(caminho_arquivo).lower().endswith(('.xlsx', '.xls'))

    if eh_excel:
        df = pd.read_excel(caminho_arquivo, sheet_name='Base_Consolidada')
    else:
        if conteudo is None:
            conteudo = Path(caminho_arquivo).read_bytes()
        df = ler_csv_flexivel(conteudo)
    if 'Code Deep' in df.columns:
        df['Code Deep'] = (
            df['Code Deep']
            .astype(str)
            .str.replace('\xa0', ' ')
            .str.strip()
        )
    validar_colunas(df, COLUNAS_OBRIGATORIAS)
    return df


# ==========================================================================
# Etapa 4 — Preparação da BASE
# ==========================================================================

def obter_coluna_com_fallback(df, nomes_aceitos, nome_padronizado, obrigatoria=True):
    for nome in nomes_aceitos:
        if nome in df.columns:
            return df[nome]
    if obrigatoria:
        raise ErroProcessamento(
            f"ERRO: coluna obrigatória não encontrada para '{nome_padronizado}'. "
            f"Nomes procurados: {nomes_aceitos}"
        )
    print(f"Aviso: coluna '{nome_padronizado}' não encontrada.")
    return None


def derivar_code(code_deep: pd.Series) -> pd.Series:
    valores = code_deep.fillna('').astype(str).str.strip()
    invalidos = valores.eq('') | valores.str.casefold().eq('nan')
    sem_ultimo = valores.str.rsplit('-', n=1).str[0]
    return pd.Series(np.where(invalidos, '', sem_ultimo), index=code_deep.index)


def preparar_base(df: pd.DataFrame) -> pd.DataFrame:
    col_situacao_origem = obter_coluna_com_fallback(
        df, COLUNAS_SITUACAO_BACKOFFICE, 'Situação Backoffice', obrigatoria=False
    )
    df['Situação'] = col_situacao_origem if col_situacao_origem is not None else ''

    df['CODE'] = derivar_code(df['Code Deep'])

    df['Alguém foi encontrado?_field'] = obter_coluna_com_fallback(
        df, COLUNAS_ALGUEM_ENCONTRADO, 'Alguém foi encontrado?_field', obrigatoria=True
    )
    df['1.4 Município_field'] = obter_coluna_com_fallback(
        df, COLUNAS_MUNICIPIO, '1.4 Município_field', obrigatoria=True
    )
    codigo_unico = obter_coluna_com_fallback(
        df, COLUNAS_CODIGO_UNICO_CLIENTE, 'codigo_unico_cliente', obrigatoria=False
    )
    if codigo_unico is not None:
        df['codigo_unico_cliente'] = codigo_unico

    col_encontrado = 'Alguém foi encontrado?_field'
    if col_encontrado in df.columns:
        respostas = df[col_encontrado].fillna('').astype(str).str.strip()
        df['Tipo_Visita'] = np.select(
            [respostas.eq('Sim'), respostas.isin(['Sim, mas se recusou a responder', 'Abandonado'])],
            ['Produtiva', 'Impedida'],
            default='Revisita'
        )
    else:
        df['Tipo_Visita'] = ''

    df['Sinalização Código'] = df.get('Status_Validacao', '')
    df['Indicativo_Visita'] = df.get('Status_Aprovacao_Backoffice', '')

    for col in COLUNAS_VAZIAS_BACKOFFICE:
        df[col] = ""

    df_base = df[[c for c in ORDEM_FINAL_BASE if c in df.columns]].copy()
    return df_base


def filtrar_pendentes(df_base: pd.DataFrame) -> pd.DataFrame:
    mask = df_base['Situação'].fillna('').astype(str).str.strip().str.casefold().eq('pendente')
    esperado = int(mask.sum())
    df_filtrado = df_base[mask].copy()
    if len(df_filtrado) != esperado:
        raise ErroProcessamento(
            f"ERRO: falha ao filtrar Pendentes (esperado {esperado}, obtido {len(df_filtrado)})."
        )
    return df_filtrado


def filtrar_por_municipio(df_base: pd.DataFrame, municipio: str) -> pd.DataFrame:
    nome_coluna = next((n for n in COLUNAS_MUNICIPIO if n in df_base.columns), None)
    if nome_coluna is None:
        raise ErroProcessamento(
            f"ERRO: coluna de município não encontrada. Procurados: {COLUNAS_MUNICIPIO}"
        )

    serie_original = df_base[nome_coluna].fillna('').astype(str).str.strip()
    serie_normalizada = serie_original.apply(normalizar_texto)

    mapa = pd.DataFrame({'original': serie_original, 'normalizado': serie_normalizada})
    mapa = mapa[mapa['normalizado'] != '']
    normalizado_para_original = mapa.groupby('normalizado')['original'].first().to_dict()
    contagem = mapa.groupby('normalizado').size().to_dict()
    valores = list(normalizado_para_original.keys())

    municipio_norm = normalizar_texto(municipio)
    aceitos = set()
    if municipio_norm in normalizado_para_original:
        aceitos.add(municipio_norm)

    similares = difflib.get_close_matches(municipio_norm, valores, n=len(valores), cutoff=LIMIAR_SIMILARIDADE_MUNICIPIO)
    aceitos.update(similares)

    for alias_list in ALIASES_MUNICIPIOS.values():
        for alias in alias_list:
            alias_norm = normalizar_texto(alias)
            if alias_norm in normalizado_para_original:
                aceitos.add(alias_norm)

    print("\nFiltro de município (fuzzy match):")
    print("Municípios ACEITOS:")
    for v in sorted(aceitos, key=lambda x: normalizado_para_original[x]):
        print(f"  - {normalizado_para_original[v]}: {contagem[v]} registro(s)")
    print("Municípios EXCLUÍDOS:")
    for v in sorted(set(valores) - aceitos, key=lambda x: normalizado_para_original[x]):
        print(f"  - {normalizado_para_original[v]}: {contagem[v]} registro(s)")
    print()

    mask_aceito = serie_normalizada.isin(aceitos)
    df_filtrado = df_base[mask_aceito].copy()
    if df_filtrado.empty:
        raise ErroProcessamento(
            f"ERRO: nenhum registro para '{municipio}'. Municípios encontrados: "
            f"{', '.join(normalizado_para_original[v] for v in valores)}"
        )
    return df_filtrado


def montar_lista_agentes(df_base: pd.DataFrame) -> list:
    agentes = df_base['Nome do Agente'].dropna().astype(str).str.strip()
    agentes = agentes[agentes != ""]
    return sorted(agentes.unique().tolist())


def listar_municipios_pendentes(arquivo_bytes: bytes) -> list[str]:
    """Lista os municípios existentes entre os registros pendentes da base."""
    dados = carregar_base(BytesIO(arquivo_bytes))
    col_situacao = next((c for c in COLUNAS_SITUACAO_BACKOFFICE if c in dados.columns), None)
    col_municipio = next((c for c in COLUNAS_MUNICIPIO if c in dados.columns), None)

    if col_situacao is None:
        raise ErroProcessamento(
            f"ERRO: coluna de situação do Backoffice não encontrada. Procuradas: {COLUNAS_SITUACAO_BACKOFFICE}"
        )
    if col_municipio is None:
        raise ErroProcessamento(
            f"ERRO: coluna de município não encontrada. Procuradas: {COLUNAS_MUNICIPIO}"
        )

    pendentes = dados[col_situacao].fillna('').astype(str).str.strip().str.casefold().eq('pendente')
    municipios = dados.loc[pendentes, col_municipio].dropna().astype(str).str.strip()
    municipios = municipios[(municipios != '') & ~municipios.str.casefold().eq('nan')]
    return sorted(municipios.unique().tolist(), key=normalizar_texto)


# ==========================================================================
# Utilitários de coluna
# ==========================================================================

def obter_letra_coluna(df_base: pd.DataFrame, nome_coluna: str) -> str:
    if nome_coluna not in df_base.columns:
        raise ErroProcessamento(f"ERRO: coluna '{nome_coluna}' não encontrada.")
    idx = df_base.columns.get_loc(nome_coluna) + 1
    return get_column_letter(idx)


def avisar_se_letra_divergente(nome_coluna: str, letra_real: str) -> None:
    esperada = LETRAS_ESPERADAS_BASE.get(nome_coluna)
    if esperada and letra_real != esperada:
        print(f"Aviso: coluna '{nome_coluna}' está em {letra_real} (esperado {esperada}).")


# ==========================================================================
# Aba CONT. CONF
# ==========================================================================

def criar_aba_cont_conf(wb, df_base: pd.DataFrame, ultima_linha_base: int):
    ws = wb.create_sheet('CONT. CONF')
    col_situacao1 = obter_letra_coluna(df_base, 'Situação.1')
    col_data_conf = obter_letra_coluna(df_base, 'Data de conferência')
    col_conferencia = obter_letra_coluna(df_base, 'Conferência')
    col_responsavel = obter_letra_coluna(df_base, 'Responsável')
    col_code_deep = obter_letra_coluna(df_base, 'Code Deep')

    for nome, letra in [('Situação.1', col_situacao1), ('Data de conferência', col_data_conf),
                        ('Conferência', col_conferencia), ('Responsável', col_responsavel),
                        ('Code Deep', col_code_deep)]:
        avisar_se_letra_divergente(nome, letra)

    ws['D2'] = 'DATA'
    ws['D3'] = datetime.now().date()
    ws['D3'].number_format = 'dd/mm/yyyy'

    ws['F2'] = 'NOME'
    ws['G2'] = 'APROVADO'
    ws['H2'] = 'REPROVADO'
    ws['I2'] = 'AGUARDANDO AJUSTE'
    ws['J2'] = 'TOTAL'
    ws['L2'] = 'Restantes'
    ws['M2'] = 'Percentual de conclusão'

    primeira_linha = 3
    for offset, responsavel in enumerate(RESPONSAVEIS_BACKOFFICE):
        linha = primeira_linha + offset
        ws[f'F{linha}'] = responsavel
        for col in ('G', 'H', 'I'):
            ws[f'{col}{linha}'] = (
                f"=COUNTIFS(BASE!${col_situacao1}:${col_situacao1},{col}$2,"
                f"BASE!${col_responsavel}:${col_responsavel},$F{linha},"
                f"BASE!${col_data_conf}:${col_data_conf},$D$3)"
            )
        ws[f'J{linha}'] = f"=SUM(G{linha}:I{linha})"

    ultima_linha_resp = primeira_linha + len(RESPONSAVEIS_BACKOFFICE) - 1
    linha_total = ultima_linha_resp + 1

    ws[f'F{linha_total}'] = 'Total'
    for col in ('G', 'H', 'I', 'J'):
        ws[f'{col}{linha_total}'] = f"=SUM({col}{primeira_linha}:{col}{ultima_linha_resp})"

    if ultima_linha_base < 2:
        ws['L3'] = "=0"
        ws['M3'] = "=0"
    else:
        ws['L3'] = f"=COUNTBLANK(BASE!{col_conferencia}2:{col_conferencia}{ultima_linha_base})"
        ws['M3'] = (
            f'=IFERROR((COUNTIF(BASE!{col_conferencia}:{col_conferencia},"FEITO")'
            f'+COUNTIF(BASE!{col_conferencia}:{col_conferencia},"ENVIAR P AGENTE"))'
            f'/COUNTA(BASE!{col_code_deep}2:{col_code_deep}{ultima_linha_base}),0)'
        )
    ws['M3'].number_format = '0.0%'

    info = {
        'primeira_linha_resp': primeira_linha,
        'ultima_linha_resp': ultima_linha_resp,
        'linha_total': linha_total,
    }
    return ws, info


def formatar_aba_cont_conf(ws: Worksheet, info: dict) -> None:
    for cel in ('D2', 'F2', 'G2', 'H2', 'I2', 'J2'):
        ws[cel].font = FONTE_HEADER_PRETA
        ws[cel].fill = FUNDO_HEADER_AZUL_CLARO
        ws[cel].border = BORDA_FINA
        ws[cel].alignment = CENTRO

    for cel in ('L2', 'M2'):
        ws[cel].font = FONTE_HEADER_BRANCA
        ws[cel].fill = FUNDO_PAINEL_VERMELHO
        ws[cel].border = BORDA_FINA
        ws[cel].alignment = CENTRO

    for cel in ('L3', 'M3'):
        ws[cel].border = BORDA_FINA
        ws[cel].alignment = CENTRO

    primeira = info['primeira_linha_resp']
    ultima = info['ultima_linha_resp']
    linha_total = info['linha_total']

    for linha in range(primeira, ultima + 1):
        for col in ('F', 'G', 'H', 'I', 'J'):
            ws[f'{col}{linha}'].border = BORDA_FINA
            ws[f'{col}{linha}'].alignment = ESQUERDA if col == 'F' else CENTRO

    for col in ('F', 'G', 'H', 'I', 'J'):
        ws[f'{col}{linha_total}'].border = BORDA_FINA
        ws[f'{col}{linha_total}'].font = FONTE_HEADER_PRETA
        ws[f'{col}{linha_total}'].fill = FUNDO_TOTAL
        ws[f'{col}{linha_total}'].alignment = ESQUERDA if col == 'F' else CENTRO

    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['F'].width = 28
    for col in ('G', 'H', 'I', 'J'):
        ws.column_dimensions[col].width = 16
    ws.column_dimensions['L'].width = 12
    ws.column_dimensions['M'].width = 22


def validar_cont_conf(ws: Worksheet, info: dict, ultima_linha_base: int) -> None:
    problemas = []
    primeira = info['primeira_linha_resp']
    ultima = info['ultima_linha_resp']
    linha_total = info['linha_total']

    for cel, func in {'G': 'COUNTIFS', 'H': 'COUNTIFS', 'I': 'COUNTIFS', 'J': 'SUM'}.items():
        val = str(ws[f'{cel}{primeira}'].value or '')
        if func not in val:
            problemas.append(f"{cel}{primeira} não contém {func}.")

    for linha in range(primeira, ultima + 1):
        for col in ('G', 'H', 'I', 'J'):
            if not str(ws[f'{col}{linha}'].value or '').startswith('='):
                problemas.append(f"{col}{linha} sem fórmula.")

    if ws[f'F{linha_total}'].value != 'Total':
        problemas.append(f"F{linha_total} não é 'Total'.")

    if ultima_linha_base < 2:
        for cel in ('L3', 'M3'):
            if str(ws[cel].value or '').strip() != '=0':
                problemas.append(f"{cel} deveria ser '=0'.")
    else:
        for cel, func in {'L3': 'COUNTBLANK', 'M3': 'IFERROR'}.items():
            if func not in str(ws[cel].value or ''):
                problemas.append(f"{cel} não contém {func}.")
            if str(ultima_linha_base) not in str(ws[cel].value or ''):
                problemas.append(f"{cel} não referencia linha {ultima_linha_base}.")

    if problemas:
        raise ErroProcessamento("Falha na validação CONT. CONF:\n" + "\n".join(problemas))


# ==========================================================================
# Aba CONT.AGENTE
# ==========================================================================

def criar_aba_cont_agente(wb, agentes: list, col_nome_agente: str):
    ws = wb.create_sheet('CONT.AGENTE')
    ws['D2'] = 'NOME DO AGENTE'
    ws['E2'] = 'CONTADOR'
    ws['F2'] = 'RESPONSAVEL'

    primeira_linha = 3
    for offset, agente in enumerate(agentes):
        linha = primeira_linha + offset
        ws[f'D{linha}'] = agente
        ws[f'E{linha}'] = f"=COUNTIF(BASE!${col_nome_agente}:${col_nome_agente},D{linha})"
        ws[f'F{linha}'] = ""

    if agentes:
        ultima_linha = primeira_linha + len(agentes) - 1
        linha_total = ultima_linha + 1
        formula_total = f"=SUM(E{primeira_linha}:E{ultima_linha})"
    else:
        ultima_linha = primeira_linha - 1
        linha_total = primeira_linha
        formula_total = "=0"

    ws[f'D{linha_total}'] = 'TOTAL'
    ws[f'E{linha_total}'] = formula_total

    info = {
        'primeira_linha': primeira_linha,
        'ultima_linha': ultima_linha,
        'linha_total': linha_total,
    }
    return ws, info


def formatar_aba_cont_agente(ws: Worksheet, info: dict) -> None:
    for cel in ('D2', 'E2', 'F2'):
        ws[cel].font = FONTE_HEADER_BRANCA
        ws[cel].fill = FUNDO_HEADER_AZUL_ESCURO
        ws[cel].border = BORDA_FINA
        ws[cel].alignment = CENTRO

    primeira = info['primeira_linha']
    ultima = info['ultima_linha']
    linha_total = info['linha_total']

    maior_nome = 0
    for linha in range(primeira, ultima + 1):
        ws[f'D{linha}'].border = BORDA_FINA
        ws[f'D{linha}'].alignment = ESQUERDA
        maior_nome = max(maior_nome, len(str(ws[f'D{linha}'].value or '')))
        ws[f'E{linha}'].border = BORDA_FINA
        ws[f'E{linha}'].alignment = CENTRO
        ws[f'F{linha}'].border = BORDA_FINA
        ws[f'F{linha}'].alignment = ESQUERDA

    for col in ('D', 'E', 'F'):
        ws[f'{col}{linha_total}'].border = BORDA_FINA
        ws[f'{col}{linha_total}'].font = FONTE_HEADER_PRETA
        ws[f'{col}{linha_total}'].fill = FUNDO_TOTAL
        ws[f'{col}{linha_total}'].alignment = CENTRO if col == 'E' else ESQUERDA

    ws.column_dimensions['D'].width = max(28, maior_nome + 2)
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 22


def validar_cont_agente(ws: Worksheet, info: dict, agentes: list) -> None:
    problemas = []
    if len(agentes) != len(set(agentes)):
        problemas.append("Agentes duplicados.")
    if any(not a.strip() for a in agentes):
        problemas.append("Agente vazio.")

    primeira = info['primeira_linha']
    ultima = info['ultima_linha']
    linha_total = info['linha_total']

    for linha in range(primeira, ultima + 1):
        if not str(ws[f'E{linha}'].value or '').startswith('=COUNTIF'):
            problemas.append(f"E{linha} sem COUNTIF.")

    if ws[f'D{linha_total}'].value != 'TOTAL':
        problemas.append(f"D{linha_total} não é 'TOTAL'.")
    if linha_total != ultima + 1:
        problemas.append("TOTAL não está imediatamente após o último agente.")

    if problemas:
        raise ErroProcessamento("Falha na validação CONT.AGENTE:\n" + "\n".join(problemas))


# ==========================================================================
# Formatação e validações da BASE
# ==========================================================================

def formatar_aba_base(ws: Worksheet, df_base: pd.DataFrame) -> None:
    total_colunas = len(df_base.columns)
    for col_idx in range(1, total_colunas + 1):
        ws.cell(row=1, column=col_idx).font = FONTE_HEADER_PRETA

    ws.freeze_panes = 'A2'
    ultima_coluna = get_column_letter(total_colunas)
    ws.auto_filter.ref = f"A1:{ultima_coluna}{len(df_base) + 1}"

    for col_idx, nome_coluna in enumerate(df_base.columns, start=1):
        largura = min(max(len(str(nome_coluna)) + 2, 12), 40)
        ws.column_dimensions[get_column_letter(col_idx)].width = largura


def ocultar_colunas_base(ws: Worksheet, df_base: pd.DataFrame) -> None:
    for nome_coluna in COLUNAS_OCULTAS_BASE:
        if nome_coluna not in df_base.columns:
            continue
        letra = obter_letra_coluna(df_base, nome_coluna)
        ws.column_dimensions[letra].hidden = True


NOME_DEFINIDO_RESPONSAVEIS = "ResponsaveisBackofficeConfCont"


def _criar_dv_lista_estatica(opcoes: list) -> DataValidation:
    formula = '"' + ",".join(opcoes) + '"'
    return DataValidation(
        type="list",
        formula1=formula,
        allow_blank=True,
        showErrorMessage=True,
        showDropDown=False,
        errorStyle="stop",
        errorTitle="Valor inválido",
        error=f"Selecione um dos valores: {', '.join(opcoes)}.",
    )


def aplicar_validacoes_base(wb, ws_base: Worksheet, df_base: pd.DataFrame,
                            ultima_linha_base: int, ultima_linha_resp: int) -> None:
    if ultima_linha_base < 2:
        return

    col_situacao1 = obter_letra_coluna(df_base, 'Situação.1')
    col_conferencia = obter_letra_coluna(df_base, 'Conferência')
    col_responsavel = obter_letra_coluna(df_base, 'Responsável')

    dv1 = _criar_dv_lista_estatica(OPCOES_SITUACAO_1)
    dv1.add(f"{col_situacao1}2:{col_situacao1}{ultima_linha_base}")
    ws_base.add_data_validation(dv1)

    dv2 = _criar_dv_lista_estatica(OPCOES_CONFERENCIA)
    dv2.add(f"{col_conferencia}2:{col_conferencia}{ultima_linha_base}")
    ws_base.add_data_validation(dv2)

    wb.defined_names[NOME_DEFINIDO_RESPONSAVEIS] = DefinedName(
        NOME_DEFINIDO_RESPONSAVEIS,
        attr_text=f"'CONT. CONF'!$F$3:$F${ultima_linha_resp}",
    )

    dv3 = DataValidation(
        type="list",
        formula1=f"={NOME_DEFINIDO_RESPONSAVEIS}",
        allow_blank=True,
        showErrorMessage=True,
        showDropDown=False,
        errorStyle="stop",
        errorTitle="Valor inválido",
        error="Selecione um responsável da lista (aba CONT. CONF, coluna F).",
    )
    dv3.add(f"{col_responsavel}2:{col_responsavel}{ultima_linha_base}")
    ws_base.add_data_validation(dv3)


def configurar_recalculo_automatico(wb) -> None:
    try:
        wb.calculation.calcMode = "auto"
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    except Exception:
        pass


# ==========================================================================
# Geração do arquivo final
# ==========================================================================

def gerar_arquivo_auditoria(df_base: pd.DataFrame, agentes: list,
                            pasta_destino: Path, nome_arquivo: str) -> Path:
    pasta_destino.mkdir(parents=True, exist_ok=True)
    caminho = pasta_destino / nome_arquivo

    with pd.ExcelWriter(caminho, engine='openpyxl') as writer:
        df_base.to_excel(writer, sheet_name='BASE', index=False)
        wb = writer.book
        ws_base = writer.sheets['BASE']
        formatar_aba_base(ws_base, df_base)
        ocultar_colunas_base(ws_base, df_base)

        ultima_linha_base = len(df_base) + 1

        ws_cont_conf, info_cont_conf = criar_aba_cont_conf(wb, df_base, ultima_linha_base)
        formatar_aba_cont_conf(ws_cont_conf, info_cont_conf)
        validar_cont_conf(ws_cont_conf, info_cont_conf, ultima_linha_base)

        aplicar_validacoes_base(
            wb, ws_base, df_base, ultima_linha_base, info_cont_conf['ultima_linha_resp']
        )

        col_nome_agente = obter_letra_coluna(df_base, 'Nome do Agente')
        ws_cont_agente, info_cont_agente = criar_aba_cont_agente(wb, agentes, col_nome_agente)
        formatar_aba_cont_agente(ws_cont_agente, info_cont_agente)
        validar_cont_agente(ws_cont_agente, info_cont_agente, agentes)

        configurar_recalculo_automatico(wb)

    print(f"✅ Arquivo de auditoria gerado com sucesso em: {caminho}")
    return caminho


def gerar_arquivo_auditoria_memoria(df_base: pd.DataFrame, agentes: list) -> bytes:
    """Gera o mesmo arquivo final em memória para download no Streamlit."""
    saida = BytesIO()
    with pd.ExcelWriter(saida, engine='openpyxl') as writer:
        df_base.to_excel(writer, sheet_name='BASE', index=False)
        wb = writer.book
        ws_base = writer.sheets['BASE']
        formatar_aba_base(ws_base, df_base)
        ocultar_colunas_base(ws_base, df_base)

        ultima_linha_base = len(df_base) + 1
        ws_cont_conf, info_cont_conf = criar_aba_cont_conf(wb, df_base, ultima_linha_base)
        formatar_aba_cont_conf(ws_cont_conf, info_cont_conf)
        validar_cont_conf(ws_cont_conf, info_cont_conf, ultima_linha_base)
        aplicar_validacoes_base(
            wb, ws_base, df_base, ultima_linha_base, info_cont_conf['ultima_linha_resp']
        )

        col_nome_agente = obter_letra_coluna(df_base, 'Nome do Agente')
        ws_cont_agente, info_cont_agente = criar_aba_cont_agente(wb, agentes, col_nome_agente)
        formatar_aba_cont_agente(ws_cont_agente, info_cont_agente)
        validar_cont_agente(ws_cont_agente, info_cont_agente, agentes)
        configurar_recalculo_automatico(wb)

    saida.seek(0)
    return saida.getvalue()


def processar_conferencia(payload, progresso=None):
    """Executa a conferência para um arquivo enviado e um município informado."""
    arquivo_bytes, municipio = payload
    municipio = normalizar_municipio(municipio)
    if not municipio:
        raise ErroProcessamento("ERRO: nome de município inválido.")

    def atualizar(percentual, etapa, detalhe=""):
        if progresso:
            progresso(percentual, etapa, detalhe)

    atualizar(0.08, "Carregando a base", "Lendo a aba Base_Consolidada")
    df_bruto = carregar_base(BytesIO(arquivo_bytes))
    atualizar(0.24, "Preparando os dados", f"{len(df_bruto):,} registros encontrados")
    df_preparado = preparar_base(df_bruto)
    atualizar(0.42, "Filtrando pendências", "Selecionando registros com situação Pendente")
    df_pendentes = filtrar_pendentes(df_preparado)
    atualizar(0.58, "Filtrando o município", municipio)
    df_filtrado = filtrar_por_municipio(df_pendentes, municipio)
    agentes = montar_lista_agentes(df_filtrado)
    atualizar(0.78, "Montando o Excel", "Criando controles, fórmulas e validações")
    excel = gerar_arquivo_auditoria_memoria(df_filtrado, agentes)
    atualizar(1.0, "Processamento concluído", "Arquivo de conferência pronto para download")
    return df_filtrado, excel


# ==========================================================================
# Execução Principal
# ==========================================================================

def main():
    try:
        municipio = solicitar_municipio()
        arquivo_input = obter_arquivo_input(PASTA_INPUT)

        print(f"\nCarregando dados de: {arquivo_input.name}...")
        df_bruto = carregar_base(arquivo_input)

        df_preparado = preparar_base(df_bruto)
        df_pendentes = filtrar_pendentes(df_preparado)
        df_filtrado = filtrar_por_municipio(df_pendentes, municipio)

        agentes = montar_lista_agentes(df_filtrado)

        data_hoje = datetime.now().strftime("%Y%m%d")
        nome_arquivo = f"Conferencia_{municipio}_{data_hoje}.xlsx"
        gerar_arquivo_auditoria(df_filtrado, agentes, PASTA_EXPORT, nome_arquivo)

        print("\n✅ Processamento concluído com sucesso!")

    except ErroProcessamento as e:
        print(f"\n❌ {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()
