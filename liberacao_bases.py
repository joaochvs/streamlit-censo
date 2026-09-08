from pathlib import Path
from copy import copy
from io import BytesIO
import sys
import pandas as pd
from openpyxl import Workbook, load_workbook

from leitura_csv import ler_csv_flexivel

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(errors="replace")

# ============================================================
# CONFIGURAÇÃO
# ============================================================

ARQUIVO_BASE = Path(r"C:\Users\LucasAragãoPereira\OneDrive - DEEPESSOAS\Área de Trabalho\AUTO_DESM\Relatorio_Master_V12.1_10h50m01s.xlsx")
ARQUIVO_TEMPLATE = Path(r"C:\Users\LucasAragãoPereira\OneDrive - DEEPESSOAS\Área de Trabalho\AUTO_DESM\template_desm.xlsx")
ARQUIVO_SAIDA = Path("BASE_LIBERACAO_PREENCHIDA.xlsx")

ABA_ORIGEM = "Base_Consolidada"
ABA_BASE = "BASE"
ABA_DESMEMBRAMENTO = "BASE DESMEMBRAMENTO"
ABA_CONTROLE = "CONT. CONF"

# Status originais do Relatorio_Master -> status que irão para o template.
STATUS_BASE = {
    "🚨 DUPLICATA - MESMO IMÓVEL EM OUTRO CÓDIGO": "Duplicata",
    "👀 REVISÃO - FOTO PARECIDA": "Revisão",
}

STATUS_DESMEMBRAMENTO = {
    "🚨 DESMEMBRAMENTO - FOTO INTRUSA": "DESMEMBRAMENTO",
}

# De/Para: coluna do template -> coluna do Relatorio_Master.
DE_PARA = {
    "A": "1.4 Município",
    "B": "Link Backoffice",
    "C": "ID_Completo",
    "E": "Status_Validacao",
    "G": "Codigo_Suspeito",
    "I": "Detalhe_Inconsistencia",
    "M": "Percentual_Similaridade",
    "N": "Apontamento_Desmembramento",
}


def cabecalhos_para_indices(ws):
    """Retorna {nome_do_cabecalho: numero_da_coluna} para a linha 1."""
    return {
        cell.value: cell.column
        for cell in ws[1]
        if cell.value is not None
    }


def validar_colunas_origem(ws):
    """Garante que todas as colunas necessárias existem no relatório."""

    print("🔎 Validando colunas da base de origem...")

    headers = cabecalhos_para_indices(ws)

    necessarias = set(DE_PARA.values()) | {"Status_Validacao"}

    faltantes = sorted(
        c for c in necessarias
        if c not in headers
    )

    if faltantes:
        raise ValueError(
            "As seguintes colunas não foram encontradas na aba "
            f"'{ABA_ORIGEM}': {', '.join(faltantes)}"
        )

    print(f"✅ Todas as {len(necessarias)} colunas necessárias foram encontradas.")

    return headers


def copiar_formatacao_linha_modelo(ws, linha_destino, linha_modelo=2):
    """
    Copia apenas a formatação da linha modelo do template.
    """

    if linha_destino == linha_modelo:
        return

    for col in range(1, ws.max_column + 1):
        origem = ws.cell(linha_modelo, col)
        destino = ws.cell(linha_destino, col)

        if origem.has_style:
            destino._style = copy(origem._style)

        if origem.number_format:
            destino.number_format = origem.number_format

        if origem.font:
            destino.font = copy(origem.font)

        if origem.fill:
            destino.fill = copy(origem.fill)

        if origem.border:
            destino.border = copy(origem.border)

        if origem.alignment:
            destino.alignment = copy(origem.alignment)

        if origem.protection:
            destino.protection = copy(origem.protection)

    if ws.row_dimensions[linha_modelo].height is not None:
        ws.row_dimensions[linha_destino].height = (
            ws.row_dimensions[linha_modelo].height
        )


def preparar_area(ws, quantidade_registros):
    """
    Preserva a formatação do template e limpa conteúdos antigos.
    """

    ultima_linha_necessaria = max(
        2,
        quantidade_registros + 1
    )

    ultima_linha_limpeza = max(
        ws.max_row,
        ultima_linha_necessaria
    )

    # Garante formatação nas linhas extras.
    for linha in range(
        ws.max_row + 1,
        ultima_linha_necessaria + 1
    ):
        copiar_formatacao_linha_modelo(ws, linha)

    # Limpa somente valores/fórmulas.
    for row in ws.iter_rows(
        min_row=2,
        max_row=ultima_linha_limpeza,
        min_col=1,
        max_col=14,
    ):
        for cell in row:
            cell.value = None


def ler_registros(ws, headers):
    """
    Faz o de/para de todos os registros usando leitura sequencial.
    Muito mais rápida que acessar ws.cell() repetidamente.
    """
    print("📖 Iniciando leitura dos registros da origem...")

    registros = []

    total_linhas = ws.max_row - 1

    print(f"📊 Linhas encontradas na origem: {total_linhas:,}")
    print("🚀 Iniciando leitura otimizada...")

    # Converte os números das colunas para índices começando em 0.
    indices = {
        col_template: headers[col_origem] - 1
        for col_template, col_origem in DE_PARA.items()
    }

    indice_id = headers["ID_Completo"] - 1

    for contador, linha in enumerate(
        ws.iter_rows(
            min_row=2,
            values_only=True
        ),
        start=1
    ):
        # ID_Completo
        id_completo = linha[indice_id]

        if id_completo in (None, ""):
            continue

        registro = {}

        for col_template, indice in indices.items():
            registro[col_template] = linha[indice]

        registros.append(registro)

        # Mostra progresso a cada 500 linhas
        if contador % 500 == 0 or contador == total_linhas:
            percentual = contador / total_linhas * 100

            print(
                f"   🔄 Linha {contador:,}/{total_linhas:,} "
                f"({percentual:.1f}%)"
            )

    print(
        f"✅ Leitura concluída! "
        f"{len(registros):,} registros válidos encontrados."
    )

    return registros


def preencher_url_backoffice_cs(registros):
    """
    Equivale ao PROCX:
    =PROCX(G2;C:C;B:B;"")&""
    """

    print("🔗 Criando índice de links do Backoffice...")

    link_por_id = {}

    for registro in registros:
        id_completo = registro.get("C")
        link = registro.get("B")

        if id_completo not in (None, ""):
            link_por_id[
                str(id_completo).strip()
            ] = link

    print(
        f"✅ Índice criado com {len(link_por_id):,} IDs."
    )

    print("🔎 Procurando links dos códigos suspeitos...")

    encontrados = 0
    sem_link = 0

    for registro in registros:

        codigo_suspeito = registro.get("G")

        if codigo_suspeito in (None, ""):
            registro["F"] = ""
            continue

        link = link_por_id.get(
            str(codigo_suspeito).strip(),
            ""
        )

        registro["F"] = "" if link is None else link

        if link:
            encontrados += 1
        else:
            sem_link += 1

    print(
        f"✅ Busca de links concluída!"
    )
    print(f"   🔗 Links encontrados: {encontrados:,}")
    print(f"   ⚠️ Links não encontrados: {sem_link:,}")


def separar_registros(registros):
    """
    Separa BASE e BASE DESMEMBRAMENTO.
    """

    print("✂️ Separando os registros por status...")

    base = []
    desmembramento = []

    for registro_original in registros:

        status_original = registro_original.get("E")

        if status_original in STATUS_BASE:

            registro = registro_original.copy()

            registro["E"] = STATUS_BASE[
                status_original
            ]

            base.append(registro)

        elif status_original in STATUS_DESMEMBRAMENTO:

            registro = registro_original.copy()

            registro["E"] = STATUS_DESMEMBRAMENTO[
                status_original
            ]

            desmembramento.append(registro)

    print("✅ Separação concluída!")
    print(f"   📋 BASE: {len(base):,} registros")
    print(
        f"   ✂️ BASE DESMEMBRAMENTO: "
        f"{len(desmembramento):,} registros"
    )

    return base, desmembramento


def gravar_registros(ws, registros):
    """
    Grava os registros no template.
    """

    print(
        f"💾 Preparando aba '{ws.title}' "
        f"para {len(registros):,} registros..."
    )

    preparar_area(ws, len(registros))

    total = len(registros)

    for contador, (linha_excel, registro) in enumerate(
        enumerate(registros, start=2),
        start=1
    ):

        # De/para / valores
        ws[f"A{linha_excel}"] = registro.get("A")
        ws[f"B{linha_excel}"] = registro.get("B")
        ws[f"C{linha_excel}"] = registro.get("C")
        ws[f"E{linha_excel}"] = registro.get("E")
        ws[f"F{linha_excel}"] = registro.get("F", "")
        ws[f"G{linha_excel}"] = registro.get("G")
        ws[f"I{linha_excel}"] = registro.get("I")
        ws[f"M{linha_excel}"] = registro.get("M")
        ws[f"N{linha_excel}"] = registro.get("N")

        # Fórmulas
        ws[f"D{linha_excel}"] = (
            f'=COUNTIF(C:C,C{linha_excel})'
        )

        ws[f"H{linha_excel}"] = (
            f'=COUNTIF(G:G,G{linha_excel})'
        )

        # Conferência
        ws[f"J{linha_excel}"] = None
        ws[f"K{linha_excel}"] = None
        ws[f"L{linha_excel}"] = None

        # Progresso a cada 5.000 registros.
        if contador % 5000 == 0 or contador == total:
            percentual = (
                contador / total * 100
                if total
                else 100
            )

            print(
                f"   📝 {ws.title}: "
                f"{contador:,}/{total:,} "
                f"({percentual:.1f}%)"
            )

    print(
        f"✅ Aba '{ws.title}' preenchida com sucesso!"
    )


def atualizar_controle(
    ws_controle,
    qtd_base,
    qtd_desmembramento
):
    """
    Ajusta automaticamente os contadores.
    """

    print("📊 Atualizando aba CONT. CONF...")

    if qtd_base > 0:

        ultima_base = qtd_base + 1

        ws_controle["N3"] = (
            f'=COUNTBLANK(BASE!J2:J{ultima_base})'
        )

        ws_controle["O3"] = (
            f'=(COUNTIF(BASE!J2:J{ultima_base},"FEITO")+'
            f'COUNTIF(BASE!J2:J{ultima_base},"MANTIDO"))/'
            f'COUNTA(BASE!C2:C{ultima_base})'
        )

    else:

        ws_controle["N3"] = 0
        ws_controle["O3"] = 0

    if qtd_desmembramento > 0:

        ultima_desm = qtd_desmembramento + 1

        ws_controle["R3"] = (
            f"=COUNTBLANK('BASE DESMEMBRAMENTO'!"
            f"J2:J{ultima_desm})"
        )

        ws_controle["S3"] = (
            f"=(COUNTIF('BASE DESMEMBRAMENTO'!"
            f"J2:J{ultima_desm},\"FEITO\")+"
            f"COUNTIF('BASE DESMEMBRAMENTO'!"
            f"J2:J{ultima_desm},\"MANTIDO\"))/"
            f"COUNTA('BASE DESMEMBRAMENTO'!"
            f"C2:C{ultima_desm})"
        )

    else:

        ws_controle["R3"] = 0
        ws_controle["S3"] = 0

    print("✅ Aba CONT. CONF atualizada!")


def processar_liberacao(arquivo_bytes, progresso=None):
    """Executa a mesma automação usando o template incorporado ao aplicativo."""
    template = Path(__file__).resolve().parent / "resources" / "template_desm.xlsx"
    if not template.exists():
        raise FileNotFoundError("O template de liberação não foi encontrado no aplicativo.")

    def atualizar(percentual, etapa, detalhe=""):
        if progresso:
            progresso(percentual, etapa, detalhe)

    atualizar(0.05, "Abrindo os arquivos", "Carregando o relatório e o template de liberação")
    if arquivo_bytes[:4].startswith(b"PK") or arquivo_bytes[:4] == b"\xd0\xcf\x11\xe0":
        wb_origem = load_workbook(BytesIO(arquivo_bytes), data_only=False, read_only=True)
    else:
        # CSV não possui abas. Converte seu conteúdo para uma planilha
        # temporária chamada Base_Consolidada e mantém o restante da lógica.
        df_origem = ler_csv_flexivel(arquivo_bytes)
        wb_temporario = Workbook()
        ws_temporaria = wb_temporario.active
        ws_temporaria.title = ABA_ORIGEM
        ws_temporaria.append(df_origem.columns.tolist())
        for linha in df_origem.itertuples(index=False, name=None):
            ws_temporaria.append([None if pd.isna(valor) else valor for valor in linha])
        del df_origem
        wb_origem = wb_temporario
    wb_saida = load_workbook(template, data_only=False)

    try:
        ws_origem = wb_origem[ABA_ORIGEM]
        ws_base = wb_saida[ABA_BASE]
        ws_desm = wb_saida[ABA_DESMEMBRAMENTO]
        ws_controle = wb_saida[ABA_CONTROLE]

        atualizar(0.15, "Validando a estrutura", f"Conferindo a aba {ABA_ORIGEM}")
        headers = validar_colunas_origem(ws_origem)

        atualizar(0.25, "Lendo os registros", f"{max(0, ws_origem.max_row - 1):,} linhas encontradas")
        registros = ler_registros(ws_origem, headers)

        atualizar(0.48, "Preenchendo os links", "Relacionando os códigos suspeitos ao Backoffice")
        preencher_url_backoffice_cs(registros)

        atualizar(0.58, "Separando os registros", "Classificando revisões, duplicatas e desmembramentos")
        registros_base, registros_desm = separar_registros(registros)

        atualizar(0.70, "Montando o arquivo", "Preenchendo a aba BASE")
        gravar_registros(ws_base, registros_base)
        atualizar(0.82, "Montando o arquivo", "Preenchendo a aba BASE DESMEMBRAMENTO")
        gravar_registros(ws_desm, registros_desm)
        atualizar_controle(ws_controle, len(registros_base), len(registros_desm))

        atualizar(0.94, "Salvando o resultado", "Gerando o Excel final")
        saida = BytesIO()
        wb_saida.save(saida)
        saida.seek(0)
        atualizar(1.0, "Processamento concluído", "Arquivo de liberação pronto para download")
        return {
            "registros": len(registros),
            "base": len(registros_base),
            "desmembramentos": len(registros_desm),
        }, saida.getvalue()
    finally:
        wb_origem.close()
        wb_saida.close()


def processar():

    print("\n" + "=" * 60)
    print("🚀 INICIANDO AUTOMAÇÃO DE LIBERAÇÃO")
    print("=" * 60)

    # --------------------------------------------------------
    # 1. Verificação dos arquivos
    # --------------------------------------------------------

    print("\n📂 ETAPA 1/7 — Verificando arquivos...")

    if not ARQUIVO_BASE.exists():
        raise FileNotFoundError(
            f"Base não encontrada: "
            f"{ARQUIVO_BASE.resolve()}"
        )

    print(f"   ✅ Base encontrada: {ARQUIVO_BASE}")

    if not ARQUIVO_TEMPLATE.exists():
        raise FileNotFoundError(
            f"Template não encontrado: "
            f"{ARQUIVO_TEMPLATE.resolve()}"
        )

    print(f"   ✅ Template encontrado: {ARQUIVO_TEMPLATE}")

    # --------------------------------------------------------
    # 2. Abrindo arquivos
    # --------------------------------------------------------

    print("\n📖 ETAPA 2/7 — Abrindo arquivos Excel...")
    print("   ⏳ Isso pode demorar dependendo do tamanho da base...")

    wb_origem = load_workbook(
        ARQUIVO_BASE,
        data_only=False,
        read_only=True
    )

    print("   ✅ Base de origem aberta.")

    wb_saida = load_workbook(
        ARQUIVO_TEMPLATE,
        data_only=False
    )

    print("   ✅ Template aberto.")

    try:

        ws_origem = wb_origem[ABA_ORIGEM]
        ws_base = wb_saida[ABA_BASE]
        ws_desm = wb_saida[ABA_DESMEMBRAMENTO]
        ws_controle = wb_saida[ABA_CONTROLE]

        # ----------------------------------------------------
        # 3. Validação
        # ----------------------------------------------------

        print("\n🔎 ETAPA 3/7 — Validando estrutura da base...")

        headers = validar_colunas_origem(
            ws_origem
        )

        # ----------------------------------------------------
        # 4. Leitura
        # ----------------------------------------------------

        print("\n📚 ETAPA 4/7 — Lendo registros...")

        registros = ler_registros(
            ws_origem,
            headers
        )

        # ----------------------------------------------------
        # 5. PROCX
        # ----------------------------------------------------

        print("\n🔗 ETAPA 5/7 — Preenchendo links do Backoffice...")

        preencher_url_backoffice_cs(
            registros
        )

        # ----------------------------------------------------
        # 6. Separação
        # ----------------------------------------------------

        print("\n✂️ ETAPA 6/7 — Separando registros...")

        registros_base, registros_desm = (
            separar_registros(registros)
        )

        # ----------------------------------------------------
        # 7. Gravação
        # ----------------------------------------------------

        print("\n💾 ETAPA 7/7 — Gravando resultado final...")

        print("\n📋 Gravando BASE...")
        gravar_registros(
            ws_base,
            registros_base
        )

        print("\n✂️ Gravando BASE DESMEMBRAMENTO...")
        gravar_registros(
            ws_desm,
            registros_desm
        )

        print("\n📊 Atualizando controles...")

        atualizar_controle(
            ws_controle,
            qtd_base=len(registros_base),
            qtd_desmembramento=len(registros_desm)
        )

        # Aba DESMEMBRADO não é alterada.

        print("\n💾 Salvando arquivo final...")
        print("   ⏳ Aguarde... o Excel está sendo gravado.")

        wb_saida.save(
            ARQUIVO_SAIDA
        )

        # ----------------------------------------------------
        # FINAL
        # ----------------------------------------------------

        print("\n" + "=" * 60)
        print("🎉 PROCESSO CONCLUÍDO COM SUCESSO!")
        print("=" * 60)

        print(
            f"📊 Registros lidos da origem: "
            f"{len(registros):,}"
        )

        print(
            f"📋 BASE - Duplicata/Revisão: "
            f"{len(registros_base):,}"
        )

        print(
            f"✂️ BASE DESMEMBRAMENTO: "
            f"{len(registros_desm):,}"
        )

        print(
            f"📁 Arquivo gerado: "
            f"{ARQUIVO_SAIDA.resolve()}"
        )

        print("=" * 60)

    finally:

        print("\n🔒 Fechando arquivos...")

        wb_origem.close()
        wb_saida.close()

        print("✅ Arquivos fechados.")
        print("🏁 Automação finalizada!")


if __name__ == "__main__":
    processar()
